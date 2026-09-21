"""Check all declared dependencies and exercise readers with synthetic fixtures."""
import json
from pathlib import Path
import sys
import tempfile
from read_data import FORMATS, read_data


def fixtures(directory, profile):
    import sqlite3
    from zipfile import ZipFile
    root, files = Path(directory), {}
    def text(suffix, content):
        path = root / ('sample' + suffix)
        path.write_text(content, encoding='utf-8')
        files[suffix] = path
    text('.csv', 'id,value\n1,"sample, quoted"\n')
    text('.tsv', 'id\tvalue\n1\tsample\n')
    text('.json', '[{"id":1,"value":"sample"}]')
    for ext in ('.jsonl', '.ndjson'):
        text(ext, '{"id":1,"value":"sample"}\n')
    for ext in ('.txt', '.md', '.log'):
        text(ext, 'sample')
    for ext in ('.db', '.sqlite', '.sqlite3'):
        path = root / ('sample' + ext)
        connection = sqlite3.connect(path)
        connection.execute('CREATE TABLE records (id INTEGER, value TEXT)')
        connection.execute("INSERT INTO records VALUES (1, 'sample')")
        connection.commit()
        connection.close()
        files[ext] = path
    path = root / 'sample.zip'
    with ZipFile(path, 'w') as archive:
        archive.writestr('sample.txt', 'sample')
    files['.zip'] = path
    if profile == 'core':
        return files
    import openpyxl
    path = root / 'sample.xlsx'
    book = openpyxl.Workbook()
    book.active.append(['id', 'value'])
    book.active.append([1, 'sample'])
    book.save(path)
    book.close()
    files['.xlsx'] = path
    if profile == 'tabular':
        return files
    import pyarrow as pa
    import pyarrow.parquet as pq
    import pyarrow.feather as feather
    table = pa.table({'id': [1], 'value': ['sample']})
    for ext, writer in [('.parquet', pq.write_table), ('.feather', feather.write_feather)]:
        path = root / ('sample' + ext)
        writer(table, path)
        files[ext] = path
    from docx import Document
    document = Document()
    document.add_paragraph('sample')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'sample'
    path = root / 'sample.docx'
    document.save(path)
    files['.docx'] = path
    text('.html', '<html><body><p>sample</p><script>ignore()</script></body></html>')
    text('.htm', '<p>sample</p>')
    text('.xml', '<records><value>sample</value></records>')
    path = root / 'sample.ods'
    with ZipFile(path, 'w') as archive:
        archive.writestr('mimetype', 'application/vnd.oasis.opendocument.spreadsheet')
        archive.writestr('META-INF/manifest.xml', '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.spreadsheet"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/></manifest:manifest>')
        archive.writestr('content.xml', '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:spreadsheet><table:table table:name="Sheet1"><table:table-row><table:table-cell office:value-type="string"><text:p>sample</text:p></table:table-cell></table:table-row></table:table></office:spreadsheet></office:body></office:document-content>')
    files['.ods'] = path
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=200)
    font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
    page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b'BT /F1 16 Tf 20 100 Td (sample) Tj ET')
    page[NameObject('/Contents')] = writer._add_object(stream)
    path = root / 'sample.pdf'
    with path.open('wb') as handle:
        writer.write(handle)
    files['.pdf'] = path
    if profile == 'ocr':
        from PIL import Image, ImageDraw
        path = root / 'sample.png'
        image = Image.new('RGB', (400, 100), 'white')
        ImageDraw.Draw(image).text((20, 25), 'SAMPLE 123', fill='black', font_size=40)
        image.save(path)
        files['.png'] = path
    return files


def requirements(path):
    for line in path.read_text(encoding='utf-8').splitlines():
        line = line.split('#', 1)[0].strip()
        if line.startswith('-r '):
            yield from requirements(path.parent / line[3:].strip())
        elif line:
            yield line


def run(profile, language='eng'):
    from importlib.metadata import version
    result = {'python': sys.executable, 'python_version': sys.version.split()[0], 'packages': {}, 'checks': {}, 'formats': {}, 'ready': False}
    try:
        if sys.version_info < (3, 10):
            raise RuntimeError('Python 3.10 or newer is required')
        if profile != 'core':
            from packaging.requirements import Requirement
            manifest = Path(__file__).resolve().parent.parent / f'requirements-{profile}.txt'
            for line in requirements(manifest):
                requirement = Requirement(line)
                installed = version(requirement.name)
                result['packages'][requirement.name] = installed
                if installed not in requirement.specifier:
                    raise RuntimeError(f'{requirement.name}=={installed} does not satisfy {requirement.specifier}')
            import pandas as pd
            import io
            if pd.read_csv(io.StringIO('id\n1\n')).iloc[0, 0] != 1:
                raise RuntimeError('pandas CSV read check failed')
            result['checks']['pandas_csv'] = 'passed'
        if profile == 'ocr':
            import pytesseract
            languages = pytesseract.get_languages(config='')
            missing = set(language.split('+')) - set(languages)
            if missing:
                raise RuntimeError('Tesseract language data missing: ' + ','.join(sorted(missing)))
            result['ocr_languages'] = languages
            result['tesseract_version'] = str(pytesseract.get_tesseract_version())
        with tempfile.TemporaryDirectory(prefix='data-to-value-probe-') as directory:
            for suffix, path in fixtures(directory, profile).items():
                report = read_data(path, ocr=profile == 'ocr', language=language)
                content = [unit.get('records', unit.get('text')) for unit in report['units']]
                passed = report['status'] == 'read' and 'sample' in json.dumps(content).lower()
                result['formats'][suffix] = {'status': 'passed' if passed else 'failed', 'reader_status': report['status']}
                if not passed:
                    raise RuntimeError('Reader smoke check failed: ' + suffix + ': ' + json.dumps(report))
            if profile == 'ocr':
                from read_data import pdf_ocr
                if 'sample' not in pdf_ocr(Path(directory)/'sample.pdf', 0, language).lower():
                    raise RuntimeError('Rendered PDF OCR smoke check failed')
                result['checks']['rendered_pdf_ocr'] = 'passed'
        levels = {'core': 0, 'tabular': 1, 'common': 2, 'ocr': 3}
        for suffix, minimum in FORMATS.items():
            if suffix not in result['formats']:
                result['formats'][suffix] = {'status': 'reader_available_not_fixture_tested' if levels[minimum] <= levels[profile] else 'not_in_profile'}
        result['ready'] = True
    except Exception as error:
        result['error'] = type(error).__name__ + ': ' + str(error)
    return result


if __name__ == '__main__':
    print(json.dumps(run(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else 'eng'), ensure_ascii=True))
